"""Step 11 - build the two workbooks and the summary.

outputs/all-relevant-articles.xlsx  every candidate above the relevance threshold
outputs/shortlist.xlsx              the n articles to decide on (+ snowball sheet),
                                    with the add-zotero-or-not column to fill in
outputs/summary.md                  numbers; the agent then writes the analysis section

Refuses to overwrite a shortlist that already has marks unless --force is given.
"""
import numpy as np
import pandas as pd
from openpyxl import load_workbook

import _bootstrap  # noqa: F401
from modules import outputs
from modules.config import project_from_args, read_csv_rows, read_json
from modules.outputs import MARK_COLUMN


def load_judgments(path):
    frame = pd.DataFrame(read_csv_rows(path))
    if len(frame):
        frame["agent_relevance"] = frame["agent_relevance"].astype(int)
    return frame


def has_marks(path):
    if not path.exists():
        return False
    workbook = load_workbook(path, data_only=True)
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if rows and MARK_COLUMN in rows[0]:
            column = rows[0].index(MARK_COLUMN)
            if any(row[column] not in (None, "") for row in rows[1:]):
                return True
    return False


def records(frame):
    return frame.replace({np.nan: ""}).to_dict("records")


def main():
    project, args = project_from_args(__doc__, lambda p: p.add_argument(
        "--force", action="store_true", help="overwrite a shortlist that already has marks"))
    config = project.config
    link_column = config["library_link_column"]
    template = config["library_link_template"]
    n_articles = int(config["n_articles"])
    min_relevance = int(config["min_agent_relevance"])
    shortlist_path = project.outputs / "shortlist.xlsx"
    if has_marks(shortlist_path) and not args.force:
        raise SystemExit(f"{shortlist_path} already contains marks. Use --force to rebuild it "
                         "(the marks will be lost), or go on to step 12.")

    project.require(project.scored, "run 08_score_candidates.py first")
    frame = pd.read_csv(project.scored, encoding="utf-8-sig", dtype={"year": str},
                        keep_default_na=False, na_values=[""])
    frame = frame[~frame["is_seed"].astype(str).str.lower().eq("true")].copy()
    for column in ("title", "authors", "year", "journal", "abstract", "keywords_hit"):
        frame[column] = frame[column].fillna("")
    judgments = load_judgments(project.judging / "search" / "judgments.csv")
    if len(judgments):
        frame = frame.merge(judgments, on="item_id", how="left")
    else:
        print("WARNING: no agent judgments found; ranking on the computed scores alone.")
        frame["agent_relevance"], frame["agent_reason"], frame["theme"] = np.nan, "", ""
    frame["doi-link"] = frame["doi"].map(outputs.doi_link)
    frame[link_column] = [outputs.library_link(template, t, d) for t, d in zip(frame["title"], frame["doi"])]

    shortlist = outputs.allocate_shortlist(frame, n_articles, config["shortlist_mode"], min_relevance,
                                           bool(config["per_journal_redistribute"])).copy()
    shortlist.insert(0, "rank", range(1, len(shortlist) + 1))
    shortlist[MARK_COLUMN] = ""
    frame["in_shortlist"] = frame["item_id"].isin(shortlist["item_id"])
    in_pool = frame["in_pool"].astype(str).str.lower().eq("true")
    relevant = frame[in_pool | (frame["agent_relevance"] >= min_relevance)].copy()
    relevant["_journal_key"] = relevant["journal"].str.lower()
    relevant = relevant.sort_values(["_journal_key", "agent_relevance", "composite_pct", "composite"],
                                    ascending=[True, False, False, False], na_position="last")

    signal_columns = [c for c in ("coupling", "n_cites_seed", "n_shared_refs", "tfidf", "specter") if c in frame]
    all_columns = (["journal", "title", "authors", "year", "relevance_level", "agent_relevance",
                    "agent_reason", "theme", "composite_pct", "composite"] + signal_columns +
                   ["n_keyword_hits", "keywords_hit", "has_abstract", "cited_by_count", "in_shortlist",
                    "doi-link", link_column, "abstract", "item_id"])
    outputs.write_workbook(project.outputs / "all-relevant-articles.xlsx", [
        {"name": "all-relevant-articles", "columns": all_columns, "rows": records(relevant),
         "links": ("doi-link", link_column)}])

    short_columns = ["rank", MARK_COLUMN, "title", "authors", "year", "journal", "agent_relevance",
                     "agent_reason", "theme", "relevance_level", "n_cites_seed", "n_shared_refs",
                     "doi-link", link_column, "abstract", "item_id"]
    sheets = [{"name": "shortlist", "columns": short_columns, "rows": records(shortlist),
               "links": ("doi-link", link_column), "dropdown": MARK_COLUMN}]

    snow_rows = []
    if project.snowball.exists() and project.snowball.stat().st_size > 10:
        snow = pd.read_csv(project.snowball, encoding="utf-8-sig", dtype={"year": str},
                           keep_default_na=False, na_values=[""])
        for column in ("title", "authors", "year", "container", "doi", "raw_reference", "abstract"):
            snow[column] = snow[column].fillna("")
        snow_judgments = load_judgments(project.judging / "snowball" / "judgments.csv")
        if len(snow_judgments):
            snow = snow.merge(snow_judgments, on="item_id", how="left")
            snow = snow[snow["agent_relevance"] >= min_relevance]
            snow = snow.sort_values(["agent_relevance", "snowball_score"], ascending=False)
        else:
            snow["agent_relevance"], snow["agent_reason"], snow["theme"] = np.nan, "", ""
        snow = snow.head(int(config["snowball"]["n_items"])).copy()
        snow.insert(0, "rank", range(1, len(snow) + 1))
        snow[MARK_COLUMN] = ""
        snow["doi-link"] = snow["doi"].map(outputs.doi_link)
        snow[link_column] = [outputs.library_link(template, t[:200], d) for t, d in zip(snow["title"], snow["doi"])]
        snow_rows = records(snow)
        sheets.append({"name": "snowball",
                       "columns": ["rank", MARK_COLUMN, "title", "authors", "year", "container", "type",
                                   "in_journal_list", "agent_relevance", "agent_reason", "theme",
                                   "n_citing_sources", "cited_by", "doi-link", link_column,
                                   "raw_reference", "item_id"],
                       "rows": snow_rows, "links": ("doi-link", link_column), "dropdown": MARK_COLUMN})
    sheets.append({"name": "how-to", "columns": ["note"], "rows": [
        {"note": f"Choose yes or no in the '{MARK_COLUMN}' column of the 'shortlist' and 'snowball' sheets."},
        {"note": "Blank counts as no. Do not delete or rename the item_id column; rows may be sorted or filtered freely."},
        {"note": "Save and close this file, then ask the agent to run the final step."}]})
    outputs.write_workbook(shortlist_path, sheets)

    # ----- summary -----
    stats = project.update_stats(all_relevant=int(len(relevant)), shortlist=int(len(shortlist)),
                                 snowball_listed=len(snow_rows))
    scoring = stats.get("scoring", {})
    t = outputs.labels(config["summary_language"])
    settings = [
        (t["s_n"], n_articles), (t["s_m"], stats.get("n_journals", "")),
        (t["s_mode"], config["shortlist_mode"]),
        (t["s_years"], f"{config['year_from'] or t['any']} - {config['year_to'] or t['any']}"),
        (t["s_signals"], ", ".join(scoring.get("signals_used", []))),
        (t["s_keywords"], "; ".join(stats.get("keywords", []))),
    ]
    funnel = [
        (t["f_docs"], stats.get("main_docs", "")),
        (t["f_seeds"], f"{stats.get('seeds_resolved', '')} / {stats.get('seeds_unique', '')}"),
        (t["f_queries"], stats.get("queries_total", "")),
        (t["f_raw"], stats.get("raw_hits", "")),
        (t["f_unique"], stats.get("unique_articles", "")),
        (t["f_noise"], sum(stats.get(k, 0) for k in (
            "excluded_book_review", "excluded_front_matter", "excluded_correction",
            "excluded_editorial", "excluded_no_title"))),
        (t["f_dup"], stats.get("excluded_duplicate_title", 0)),
        (t["f_seed_zotero"], f"{stats.get('excluded_seed', 0)} / {stats.get('excluded_in_zotero', 0)}"),
        (t["f_scored"], stats.get("scored", "")),
        (t["f_read"], stats.get("judged_search", 0)),
        (t["f_all"], len(relevant)),
        (t["f_short"], len(shortlist)),
        (t["f_snow"], f"{stats.get('snowball_candidates', 0)} / {len(snow_rows)}"),
    ]
    journal_rows = []
    for journal, group in frame.groupby("journal"):
        journal_rows.append((journal, len(group),
                             int(group["in_pool"].astype(str).str.lower().eq("true").sum()),
                             int((group["agent_relevance"] >= min_relevance).sum()),
                             int(group["in_shortlist"].sum())))
    journal_rows.sort(key=lambda r: (-r[4], -r[3], r[0]))
    themes = shortlist["theme"].replace("", np.nan).dropna().value_counts()
    quality = [
        t["q_abstract"].format(v=scoring.get("abstract_coverage", 0)),
        t["q_refs"].format(v=scoring.get("reference_coverage", 0)),
        t["q_cites"].format(n=scoring.get("cites_a_seed", 0)),
        t["q_seeds"].format(a=stats.get("seeds_resolved", "?"), b=stats.get("seeds_unique", "?"),
                            c=stats.get("seeds_with_reference_list", "?")),
    ]
    if scoring.get("seeds_found_by_search"):
        quality.append(t["q_sanity"].format(n=scoring["seeds_found_by_search"],
                                            v=scoring.get("seeds_in_pool_share", 0)))
    if stats.get("queries_failed"):
        quality.append(t["q_failed"].format(n=stats["queries_failed"]))
    if stats.get("snowball_dois_unknown_to_crossref"):
        quality.append(t["q_unknown"].format(n=stats["snowball_dois_unknown_to_crossref"]))
    quality.append(t["q_chinese"])
    summary = outputs.build_summary(config["summary_language"], settings, funnel, journal_rows,
                                    list(themes.items()), records(shortlist.head(15)), snow_rows[:10], quality)
    (project.outputs / "summary.md").write_text(summary, encoding="utf-8")

    print(f"all-relevant-articles.xlsx  {len(relevant)} rows")
    print(f"shortlist.xlsx              {len(shortlist)} shortlist rows, {len(snow_rows)} snowball rows")
    print(f"summary.md                  written; the analysis section is still to be filled in by the agent")
    print(f"-> {project.outputs}")
    if len(shortlist) < n_articles:
        print(f"NOTE: only {len(shortlist)} of the requested {n_articles} met the minimum relevance "
              f"({min_relevance}). Lower min_agent_relevance or widen the search if more are needed.")


if __name__ == "__main__":
    main()
