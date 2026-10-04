"""Excel workbooks, the shortlist allocation, library links and the summary."""
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MARK_COLUMN = "add-zotero-or-not"
YES_VALUES = {"yes", "y", "true", "1", "v", "x", "是", "要", "✓", "✔"}
RANK_COLUMNS = ["agent_relevance", "composite_pct", "composite"]

WIDTHS = {"title": 60, "abstract": 80, "agent_reason": 60, "authors": 30, "journal": 30,
          "container": 30, "theme": 24, "keywords_hit": 40, "cited_by": 50, "raw_reference": 60,
          "doi-link": 34, "item_id": 30, MARK_COLUMN: 18, "note": 80}
WRAP = {"title", "abstract", "agent_reason", "cited_by", "raw_reference", "note"}


# ----- links -----

def library_link(template, title, doi=""):
    """Search-results URL in the library catalogue for one title."""
    if not template or not title:
        return ""
    # Primo splits its query on commas (field,precision,term), so they cannot stay in the term.
    term = " ".join(str(title).replace(",", " ").replace(";", " ").split())
    return template.replace("{title}", quote(term, safe="")).replace("{doi}", quote(doi or "", safe=""))


def doi_link(doi):
    return f"https://doi.org/{doi}" if doi else ""


# ----- shortlist -----

def rank_frame(frame):
    """Best first: the agent's rating, then the stratified composite percentile."""
    columns = [c for c in RANK_COLUMNS if c in frame]
    return frame.sort_values(columns, ascending=False, na_position="last", kind="mergesort")


def allocate_shortlist(frame, n_articles, mode, min_relevance, redistribute=True):
    """Pick the shortlist.

    global       the n best eligible articles overall
    per-journal  n // m per journal (only eligible ones); unused places go to
                 the best remaining articles overall when redistribute is on."""
    judged = "agent_relevance" in frame and frame["agent_relevance"].notna().any()
    eligible = frame[frame["agent_relevance"] >= min_relevance] if judged else frame
    ranked = rank_frame(eligible)
    if mode != "per-journal":
        return ranked.head(n_articles)
    journals = list(dict.fromkeys(frame["journal"]))
    quota = max(n_articles // max(len(journals), 1), 1)
    picked = ranked.groupby("journal", sort=False).head(quota)
    if redistribute and len(picked) < n_articles:
        rest = ranked.drop(picked.index)
        picked = pd.concat([picked, rest.head(n_articles - len(picked))])
    return rank_frame(picked).head(n_articles)


# ----- excel -----

def write_workbook(path, sheets):
    """sheets: list of dicts {name, columns, rows, links, dropdown, note}."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="DDE6F0")
    for spec in sheets:
        sheet = workbook.create_sheet(spec["name"][:31])
        columns = spec["columns"]
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in spec["rows"]:
            sheet.append([_cell_value(row.get(col)) for col in columns])
        for index, column in enumerate(columns, 1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = WIDTHS.get(column, 14)
            is_link = column in spec.get("links", ())
            for cell in sheet[letter][1:]:
                if column in WRAP:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                if is_link and cell.value:
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
        sheet.freeze_panes = "A2"
        if spec["rows"]:
            sheet.auto_filter.ref = sheet.dimensions
        dropdown = spec.get("dropdown")
        if dropdown and dropdown in columns and spec["rows"]:
            letter = get_column_letter(columns.index(dropdown) + 1)
            validation = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
            validation.error = "Please choose yes or no"
            sheet.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}{len(spec['rows']) + 1}")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _cell_value(value):
    if value is None:
        return ""
    if isinstance(value, float) and value != value:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, str) and len(value) > 32000:
        value = value[:32000]
    return value


def excel_lock_file(path):
    """Excel keeps a ~$name.xlsx file next to an open workbook."""
    path = Path(path)
    lock = path.with_name("~$" + path.name)
    return lock if lock.exists() else None


def read_marks(path, sheet_names):
    """Rows marked yes in the shortlist workbook, matched by item_id (not by position).
    Returns (selected rows, counts per sheet, unreadable mark values)."""
    workbook = load_workbook(path, data_only=True)
    selected, counts, odd = [], {}, []
    for name in sheet_names:
        if name not in workbook.sheetnames:
            continue
        rows = list(workbook[name].iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        if MARK_COLUMN not in header or "item_id" not in header:
            raise ValueError(f"sheet '{name}' has lost its '{MARK_COLUMN}' or 'item_id' column")
        marked = 0
        for values in rows[1:]:
            row = dict(zip(header, values))
            mark = str(row.get(MARK_COLUMN) or "").strip().lower()
            if not row.get("item_id"):
                continue
            if mark in YES_VALUES:
                marked += 1
                selected.append({"sheet": name, **{k: ("" if v is None else v) for k, v in row.items()}})
            elif mark and mark not in ("no", "n", "false", "0", "否", "不要"):
                odd.append(f"{name}: '{row.get(MARK_COLUMN)}' ({row['item_id']})")
        counts[name] = {"rows": len(rows) - 1, "marked_yes": marked}
    return selected, counts, odd


# ----- summary -----

TEXT = {
    "zh-TW": {
        "title": "文獻蒐集總結", "settings": "本次設定", "funnel": "數量漏斗", "journals": "各期刊結果",
        "themes": "主題分布（短名單）", "top": "最相關的文章", "snowball": "引文回溯：期刊檢索以外的重要文獻",
        "quality": "資料品質提醒", "analysis": "分析",
        "agent_todo": "<!-- AGENT: 請依短名單與候選池撰寫本節：(1) 整體觀察 (2) 與主題特別相關的文章群及其原因 (3) 這次檢索沒有涵蓋到的缺口與補救建議。完成後刪除此註解。 -->",
        "stage": "階段", "count": "數量", "journal": "期刊", "candidates": "候選", "pool": "過門檻",
        "relevant": "agent 判定相關", "short": "短名單", "theme": "主題", "why": "理由", "none": "（無）",
        "cited_by": "被 {n} 個來源引用", "any": "不限",
        "s_n": "短名單篇數 (n)", "s_m": "期刊數 (m)", "s_mode": "短名單模式", "s_years": "出版年份",
        "s_signals": "評分訊號", "s_keywords": "關鍵詞",
        "f_docs": "主要文件", "f_seeds": "種子文獻（解析出 DOI／去重後）", "f_queries": "查詢次數（關鍵詞 × 期刊）",
        "f_raw": "檢索命中", "f_unique": "不重複文章", "f_noise": "排除：書評等非研究論文",
        "f_dup": "排除：標題重複", "f_seed_zotero": "排除：本身是種子／已在 Zotero",
        "f_scored": "進入評分", "f_read": "agent 細讀", "f_all": "all-relevant-articles.xlsx",
        "f_short": "短名單", "f_snow": "引文回溯：候選／列出",
        "q_abstract": "{v:.0%} 的文章有摘要；沒有摘要的文章在自己的群組內排名，以免被系統性低估。",
        "q_refs": "{v:.0%} 的文章有公開的參考文獻清單；其餘的引文訊號是「缺少」而非 0。",
        "q_cites": "{n} 篇文章至少引用一篇種子文獻。",
        "q_seeds": "種子文獻解析出 DOI：{a} / {b}；其中有參考文獻清單的：{c}。",
        "q_sanity": "檢核：檢索本身找回了 {n} 篇種子文獻，其中 {v:.0%} 排在門檻之上。",
        "q_failed": "最近一次檢索有 {n} 次查詢失敗，重跑檢索步驟即可補上。",
        "q_unknown": "{n} 個被引用的 DOI 不在 Crossref，未列入引文回溯。",
        "q_chinese": "Crossref 幾乎不收錄中文文獻；如有需要請另外檢索。",
    },
    "en": {
        "title": "Literature collection summary", "settings": "Settings", "funnel": "Funnel", "journals": "Results by journal",
        "themes": "Themes (shortlist)", "top": "Most relevant articles", "snowball": "Snowballing: important works beyond the journal search",
        "quality": "Data-quality notes", "analysis": "Analysis",
        "agent_todo": "<!-- AGENT: write this section from the shortlist and the pool: (1) overall picture (2) clusters most relevant to the topic and why (3) gaps this search did not cover and how to fill them. Delete this comment when done. -->",
        "stage": "Stage", "count": "Count", "journal": "Journal", "candidates": "Candidates", "pool": "Above threshold",
        "relevant": "Judged relevant", "short": "Shortlist", "theme": "Theme", "why": "Reason", "none": "(none)",
        "cited_by": "cited by {n} sources", "any": "any",
        "s_n": "Shortlist size (n)", "s_m": "Journals (m)", "s_mode": "Shortlist mode", "s_years": "Publication years",
        "s_signals": "Signals", "s_keywords": "Keywords",
        "f_docs": "Main documents", "f_seeds": "Seed references (resolved to a DOI / unique)", "f_queries": "Queries (keywords x journals)",
        "f_raw": "Raw hits", "f_unique": "Unique articles", "f_noise": "Removed: book reviews and other non-articles",
        "f_dup": "Removed: duplicate titles", "f_seed_zotero": "Removed: already a seed / already in Zotero",
        "f_scored": "Scored", "f_read": "Read by the agent", "f_all": "all-relevant-articles.xlsx",
        "f_short": "Shortlist", "f_snow": "Snowballing: candidates / listed",
        "q_abstract": "{v:.0%} of the articles have an abstract; those without are ranked within their own group so they are not systematically undervalued.",
        "q_refs": "{v:.0%} have a public reference list; for the rest the citation signal is missing, not zero.",
        "q_cites": "{n} articles cite at least one seed reference.",
        "q_seeds": "Seeds resolved to a DOI: {a} / {b}; with a reference list: {c}.",
        "q_sanity": "Check: the search itself found {n} of the seeds, and {v:.0%} of those rank above the threshold.",
        "q_failed": "{n} queries failed in the last search run; re-run the search step to fill them in.",
        "q_unknown": "{n} cited DOIs are not registered with Crossref and were left out of the snowball.",
        "q_chinese": "Crossref indexes little Chinese-language scholarship; search it separately if needed.",
    },
}


def labels(language):
    return TEXT.get(language, TEXT["en"])


def _table(header, rows):
    lines = ["| " + " | ".join(header) + " |", "|" + "|".join("---" for _ in header) + "|"]
    for row in rows:
        lines.append("| " + " | ".join(str(v).replace("|", "/").replace("\n", " ") for v in row) + " |")
    return "\n".join(lines)


def build_summary(language, settings, funnel, journal_rows, theme_rows, top_rows, snowball_rows, quality):
    """Markdown summary. The numbers are written here; the agent adds the analysis."""
    t = labels(language)
    parts = [f"# {t['title']}", ""]
    parts += [f"## {t['settings']}", ""] + [f"- **{k}**: {v}" for k, v in settings] + [""]
    parts += [f"## {t['funnel']}", "", _table([t["stage"], t["count"]], funnel), ""]
    parts += [f"## {t['journals']}", "",
              _table([t["journal"], t["candidates"], t["pool"], t["relevant"], t["short"]], journal_rows), ""]
    parts += [f"## {t['themes']}", "",
              _table([t["theme"], t["count"]], theme_rows) if theme_rows else t["none"], ""]
    parts += [f"## {t['top']}", ""]
    parts += [f"{i}. **{r['title']}** — {r['authors']} ({r['year']}), *{r['journal']}*. "
              f"{t['why']}: {r['agent_reason'] or '-'}" for i, r in enumerate(top_rows, 1)] or [t["none"]]
    parts += ["", f"## {t['snowball']}", ""]
    parts += [f"{i}. **{r['title'][:160]}** — {r['authors']} ({r['year']}). "
              f"{t['cited_by'].format(n=r['n_citing_sources'])}. {r.get('agent_reason') or ''}"
              for i, r in enumerate(snowball_rows, 1)] or [t["none"]]
    parts += ["", f"## {t['quality']}", ""] + [f"- {line}" for line in quality] + [""]
    parts += [f"## {t['analysis']}", "", t["agent_todo"], ""]
    return "\n".join(parts)
